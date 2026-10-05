import os
import pandas as pd,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from cats import CATS
R=pd.read_pickle('R.pkl');M=pd.read_pickle('M.pkl');LU=pd.read_pickle('LU.pkl');SG=pd.read_pickle('SG.pkl').set_index('sku')
LAB={c[0]:c[1] for c in CATS}
df=pd.read_excel(os.environ['LK_FILE']).set_index('פריט')
def br(p): return 'SIGNET' if p<476 else ('KENDO' if 528<=p<=613 else 'מותג אחר (פלד)')
rows=[];lku=[]
pil=set(M.lk)|set(LU.lk_sku)
for r in M.itertuples():
    rows.append(dict(lk=r.lk,lkn=r.lkn,dep=df.loc[r.lk,'שם מחלקה'],sku=r.sku,desc=r.desc,conf=r.conf,note=r.notes,page=int(r.page),price=r.price,brand='SIGNET'))
for r in LU.itertuples():
    lku.append(dict(lk=r.lk_sku,lkn=r.lk_name,dep=df.loc[r.lk_sku,'שם מחלקה'],sub=r.sub,why='אין מידה/דגם תואם',note='(פיילוט) '+'אין מידה זו בעמודי סיגנט הרלוונטיים'))
for r in R.itertuples():
    if r.lk in pil: continue
    if r.status=='match':
        rows.append(dict(lk=r.lk,lkn=r.lkn,dep=r.dep,sku=r.sku,desc=r.desc,conf=r.conf,note=r.note,page=int(r.page),price=r.price,brand=r.brand))
    else:
        lku.append(dict(lk=r.lk,lkn=r.lkn,dep=r.dep,sub=r.sub,why='אין קטגוריה מקבילה' if r.status=='nocat' else 'אין מידה/דגם תואם',note=r.note))
A=pd.DataFrame(rows);U=pd.DataFrame(lku)
o={'ודאי':0,'סביר':1,'לבדיקה':2}
A=A.assign(o=A.conf.map(o)).sort_values(['o','dep','page','lk']).drop(columns='o')
U=U.sort_values(['why','dep','sub','lk'])
used=set(A.sku)
SU=SG[~SG.index.isin(used)].reset_index().sort_values(['page','sku'])
def num(p):
    try: return float(str(p).replace(',',''))
    except: return None
wb=Workbook();F='Arial';thin=Side(style='thin',color='BFBFBF');hf=PatternFill('solid',fgColor='1F3864')
fills={'ודאי':'C6EFCE','סביר':'FFEB9C','לבדיקה':'F8CBAD'}
def sheet(ws,headers,data,widths,conf_col=None,text_cols=(),money_cols=()):
    ws.sheet_view.rightToLeft=True;ws.append(headers)
    for c in ws[1]: c.font=Font(name=F,bold=True,color='FFFFFF');c.fill=hf;c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for r in data: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font=Font(name=F,size=10);c.alignment=Alignment(vertical='top',wrap_text=True,horizontal='right');c.border=Border(bottom=thin)
            if c.column in text_cols: c.number_format='@'
            if c.column in money_cols: c.number_format='#,##0.00'
        if conf_col:
            cc=row[conf_col-1];cc.fill=PatternFill('solid',fgColor=fills.get(cc.value,'FFFFFF'));cc.alignment=Alignment(horizontal='center',vertical='top')
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=32
ws=wb.active;ws.title='התאמות'
sheet(ws,['מק"ט ל.כ','תיאור ל.כ','מחלקה ל.כ','מק"ט סיגנט','תיאור סיגנט','רמת ביטחון','הערות','עמוד בקטלוג סיגנט','מחיר מחירון סיגנט (₪)','מותג בקטלוג פלד'],
 [[int(r.lk),r.lkn,r.dep,r.sku,r.desc,r.conf,r.note,r.page,num(r.price),r.brand] for r in A.itertuples()],
 [11,42,22,11,50,10,45,9,11,14],conf_col=6,text_cols=(4,),money_cols=(9,))
ws2=wb.create_sheet('ל.כ ללא התאמה')
sheet(ws2,['מק"ט ל.כ','תיאור ל.כ','מחלקה','קבוצת משנה','סיבה','הערות'],
 [[int(r.lk),r.lkn,r.dep,r.sub,r.why,r.note] for r in U.itertuples()],[11,45,24,26,18,70])
ws3=wb.create_sheet('סיגנט ללא התאמה')
sheet(ws3,['מק"ט סיגנט','תיאור סיגנט','קטגוריה','מותג בקטלוג פלד','עמוד בקטלוג','מחיר מחירון (₪)','מקור התיאור'],
 [[r.sku,r.desc,LAB.get(r.cat,''),r.brand,int(r.page),num(r.price),r.src] for r in SU.itertuples()],[11,60,22,16,9,11,12],text_cols=(1,),money_cols=(6,))
ws4=wb.create_sheet('סיכום והסבר');ws4.sheet_view.rightToLeft=True
def put(vals,bold=False,size=10):
    ws4.append(vals)
    for c in ws4[ws4.max_row]: c.font=Font(name=F,bold=bold,size=size);c.alignment=Alignment(horizontal='right',wrap_text=False)
put(['השוואת קטלוג ל.כ מול קטלוג פלד/סיגנט 2024/25 – הרצה מלאה'],True,13)
put(['מקור: https://peledtech.com/catalog2024-25/ – שכבת הטקסט של 644 העמודים ואינדקס המק"טים (עמ׳ 631–643, 4,171 מק"טים).'])
put([])
put(['מחלקה ל.כ','פריטים בקטלוג','ודאי','סביר','לבדיקה','ללא התאמה – אין קטגוריה','ללא התאמה – אין מידה/דגם'],True)
start=ws4.max_row+1
for dep in df['שם מחלקה'].value_counts().index:
    r=ws4.max_row+1;q=f'C{r}'
    ws4.append([dep,int((df['שם מחלקה']==dep).sum())]+
      [f"=COUNTIFS('התאמות'!$C:$C,$A{r},'התאמות'!$F:$F,\"{c}\")" for c in ['ודאי','סביר','לבדיקה']]+
      [f"=COUNTIFS('ל.כ ללא התאמה'!$C:$C,$A{r},'ל.כ ללא התאמה'!$E:$E,\"{w}\")" for w in ['אין קטגוריה מקבילה','אין מידה/דגם תואם']])
end=ws4.max_row;r=end+1
ws4.append(['סה"כ']+[f'=SUM({get_column_letter(c)}{start}:{get_column_letter(c)}{end})' for c in range(2,8)])
for row in ws4.iter_rows(min_row=start,max_row=r):
    for c in row: c.font=Font(name=F,size=10,bold=(c.row==r))
put([])
put(['כללי רמת ביטחון'],True,12)
for t in ['ודאי – אותו סוג מוצר, אותה מידה ואותם מאפיינים (דרייב, עומק, צלעות, סוג ראש), מתוך טבלת מידות בסיגנט, ומועמד יחיד. רק בקטגוריות עם טבלאות מובנות (בוקסות, מפתחות, אלן, ביטים, מקדחים).',
          'סביר – המידה והקטגוריה תואמות, אך פרט אחד לא מאומת: סוג/חומר לא מצוין בסיגנט, כמה דגמים באותה מידה, או קטגוריה שבה מידה לבדה לא מבטיחה דגם זהה.',
          'לבדיקה – התאמה חלקית: מידה חלקית, מספר צלעות שונה, התאמה לפי מילים בלבד, או מק"ט מתוך עמוד עם הרבה פריטים בטקסט חופשי. מוצג המועמד הטוב ביותר + חלופות.',
          'עמודות "התאמה לפי טקסט חופשי": בעמודים בלי טבלת מידות סדר הטקסט באתר מעורבב, ולכן שיוך התיאור למק"ט מקורב – כדאי לפתוח את העמוד באתר (peledtech.com/catalog2024-25/<מספר עמוד>).',
          'הפיילוט (בוקסות 1/4"–1/2", מפתחות רינג פתוח ורינג פתוח רצ׳ט) שולב כפי שאושר.',
          'מחלקות ל.כ שאין להן מקבילה בקטלוג (רתכות, ציוד הרמה, שרשראות, שינוע, 4x4, הידראוליקה) סומנו "אין קטגוריה מקבילה", חוץ מפריטים כמו כפפות, משקפיים ופנסים.',
          'קטלוג פלד כולל גם KENDO ומותגים נוספים (עמ׳ 476 ואילך); המותג מסומן לפי טווח העמודים ועשוי להיות לא מדויק בעמודי מעבר.']:
    put([t])
ws4.column_dimensions['A'].width=40
for c in 'BCDEFG': ws4.column_dimensions[c].width=14
wb.move_sheet('סיכום והסבר',offset=-3)
wb.calculation.fullCalcOnLoad=True;out=os.environ['OUT_FILE'];wb.save(out)
print(len(A),A.conf.value_counts().to_dict(),len(U),U.why.value_counts().to_dict(),len(SU))
